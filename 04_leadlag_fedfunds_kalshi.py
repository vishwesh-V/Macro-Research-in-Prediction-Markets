# /// script
# requires-python = ">=3.12"
# dependencies = ["marimo", "duckdb", "numpy", "pandas", "pyarrow", "altair", "openpyxl"]
#
# [tool.uv]
# python-preference = "only-managed"  # Anaconda's Python misreports macOS as 10.16, so uv rejects modern wheels
# ///

import marimo

__generated_with = "0.25.1"
app = marimo.App(width="medium")

with app.setup(hide_code=True):
    import calendar
    import math
    from pathlib import Path

    import altair as alt
    import duckdb
    import marimo as mo
    import numpy as np
    import pandas as pd

    ROOT = Path(__file__).resolve().parent
    FF_XLSX = ROOT / "data/KalshiProjectFedFunds.xlsx"
    # 1-minute bid/ask bars from data/KalshiProjectFFIntraday_template.xlsx, refreshed on a Bloomberg Terminal.
    INTRADAY_XLSX = ROOT / "data/bloomberg/ff_intraday.xlsx"
    INTRADAY_CACHE = ROOT / "data/bloomberg/ff_intraday_1m.parquet"  # rebuilt when the workbook is newer
    CLEAN_CANDLES = ROOT / "data/cleaned/kalshi_candles_1m.parquet"
    TZ = "America/New_York"  # the Terminal's time zone: no bars 17:00-18:00, CME's daily pause in ET
    FF_COLS = [f"FF{n}" for n in range(1, 14)]
    SAMPLE_START = pd.Timestamp("2021-01-01")
    SPREAD_WINDOW = 10  # business days in the rolling EFFR-minus-midpoint median, as in 03_fedfunds_vs_kalshi.py
    DECISION_HOUR = 14  # FOMC statements are released at 14:00 ET
    LAGS = {1: 10, 5: 6}  # lags in the Granger regressions per sampling interval: 10 and 30 minutes (even)
    EVENT_QUANTILE = 0.999  # "large move": the top 0.1% of absolute 1-minute changes
    EVENT_HORIZON = 120  # minutes after a large move in the event study

    alt.data_transformers.disable_max_rows()


@app.function(hide_code=True)
def load_futures(path):
    """Daily fed funds futures 1-13, EFFR and target range, cleaned as in 03_fedfunds_vs_kalshi.py: rows whose
    13 prices all repeat the previous row are holidays carried forward by the pull and are dropped."""
    raw = pd.read_excel(path, "Data", header=0, skiprows=[1])  # row 2 holds descriptions, not data
    raw.columns = ["date", *FF_COLS, "effr", "upper", "lower"]
    raw["date"] = pd.to_datetime(raw.date)
    raw = raw[raw.date >= SAMPLE_START - pd.DateOffset(months=1)].reset_index(drop=True)
    raw[raw.columns[1:]] = raw[raw.columns[1:]].apply(pd.to_numeric, errors="coerce")  # #N/A -> NaN
    stale = raw[FF_COLS].diff().abs().sum(axis=1).eq(0)
    return raw[~stale].reset_index(drop=True)


@app.function(hide_code=True)
def load_fomc(path):
    """FOMC decision dates from the workbook."""
    f = pd.read_excel(path, "FOMC_Dates")
    f.columns = ["meeting_date", "effective_date", "scheduled", "decision_bp", "upper_after", "lower_after", "notes"]
    f = f[f.meeting_date != "Start level"].copy()
    f["meeting_date"] = pd.to_datetime(f.meeting_date)
    return f.sort_values("meeting_date").reset_index(drop=True)


@app.function(hide_code=True)
def add_starting_rate(futures, meeting_dates):
    """Starting rate r0: the range midpoint plus the rolling median EFFR spread, as in 03_fedfunds_vs_kalshi.py."""
    midpoint = (futures.upper + futures.lower) / 2
    spread = (futures.effr - midpoint).where(~futures.date.isin(meeting_dates))
    spread_median = spread.rolling(SPREAD_WINDOW, min_periods=3).median().ffill().bfill()
    return futures.assign(r0=midpoint + spread_median)


@app.function(hide_code=True)
def load_intraday_bars(path, cache):
    """Long table of the Bloomberg 1-minute bars: contract_month, side (bid/ask), ts (end of the bar, ET), price.

    Bloomberg labels a bar by the minute it starts and Kalshi labels a candle by the minute it ends, so the
    Bloomberg label is moved one minute later: both then give the quote in force at the end of the labeled minute.
    """
    if cache.exists() and cache.stat().st_mtime > path.stat().st_mtime:
        bars = pd.read_parquet(cache)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for name in wb.sheetnames[2:]:  # README and Settings first, then one sheet per contract month
            rows = list(wb[name].iter_rows(min_row=6, max_col=5, values_only=True))
            for side, (i, j) in {"bid": (0, 1), "ask": (3, 4)}.items():
                d = pd.DataFrame([(r[i], r[j]) for r in rows if r[i] is not None], columns=["ts", "price"])
                parts.append(d.assign(contract_month=name, side=side))
        bars = pd.concat(parts, ignore_index=True)
        bars["price"] = pd.to_numeric(bars.price, errors="coerce")
        bars.dropna().to_parquet(cache)
    bars["ts"] = pd.to_datetime(bars.ts).dt.tz_localize(TZ) + pd.Timedelta(minutes=1)
    return bars


@app.function(hide_code=True)
def cme_minutes(start, end):
    """Every minute CME Globex trades fed funds futures: Sunday 18:00 to Friday 17:00 ET, paused 17:00-18:00."""
    grid = pd.date_range(start.ceil("min"), end, freq="min", tz=TZ)
    wd, hr = grid.dayofweek, grid.hour
    return grid[~((wd == 5) | ((wd == 4) & (hr >= 17)) | ((wd == 6) & (hr < 18)) | (hr == 17))]


@app.function(hide_code=True)
def next_meeting_move(prices, times, sessions, meetings, daily):
    """Futures-implied expected move (bp) at the next FOMC meeting, one value per row, by the method of
    03_fedfunds_vs_kalshi.py (the guide's Steps 1-5) applied to whatever prices the rows hold.

    prices: one column per contract month ("YYYY-MM"), quoted as 100 minus the rate.
    times: naive ET timestamps; the next meeting is the first decision (14:00 ET) after each.
    sessions: the date each row's r0 and realized EFFR are taken from.
    daily: effr and r0 by calendar date, forward-filled.

    Method B: if the month after the meeting has no meeting, its contract is the post-meeting rate. Method A
    otherwise: un-blend the meeting month, using the realized EFFR when the meeting is in the current month.
    """
    meeting_idx = pd.DatetimeIndex(meetings)
    decision = meeting_idx + pd.Timedelta(hours=DECISION_HOUR)
    D = pd.DatetimeIndex(meeting_idx[np.searchsorted(decision, times, side="right")])
    sessions = pd.DatetimeIndex(sessions)
    r_before = daily.r0.reindex(sessions).to_numpy()
    cum_effr = daily.effr.cumsum()
    with_meeting = {(d.year, d.month) for d in meetings}
    move = np.full(len(times), np.nan)
    method = np.full(len(times), "", dtype=object)
    for d in D.unique():
        sel = np.flatnonzero(D == d)
        after = d + pd.offsets.MonthBegin(1)
        after_key, own_key = f"{after.year}-{after.month:02d}", f"{d.year}-{d.month:02d}"
        if (after.year, after.month) not in with_meeting and after_key in prices:
            r_after = 100 - prices[after_key].to_numpy()[sel]
            method[sel] = "B"
        elif own_key in prices:
            n_days = calendar.monthrange(d.year, d.month)[1]
            rbar = 100 - prices[own_key].to_numpy()[sel]
            s = sessions[sel]
            same = (s.year == d.year) & (s.month == d.month)
            month_start = pd.Timestamp(d.year, d.month, 1)
            # EFFR already printed this month: days 1 .. session day - 1 (eq. 8 of the guide)
            before_start = cum_effr.get(month_start - pd.Timedelta(days=1), 0.0)
            realized = np.where(same, cum_effr.reindex(s - pd.Timedelta(days=1)).to_numpy() - before_start, 0.0)
            realized = np.where(same & (s.day == 1), 0.0, realized)
            n_real = np.where(same, s.day - 1, 0)
            r_after = (n_days * rbar - realized - (d.day - n_real) * r_before[sel]) / (n_days - d.day)
            method[sel] = "A"
        else:
            continue
        move[sel] = (r_after - r_before[sel]) * 100
    return pd.DataFrame({"meeting": D, "method": method, "futures_bp": move})


@app.function(hide_code=True)
def chi2_sf_even(w, df):
    """P(chi-squared with an even number df of degrees of freedom > w), in closed form."""
    half = w / 2
    return math.exp(-half) * sum(half**i / math.factorial(i) for i in range(df // 2))


@app.function(hide_code=True)
def granger(changes, lags):
    """Granger tests in both directions on changes within a session.

    For each market, its change is regressed on `lags` lags of its own changes and of the other market's,
    by OLS with standard errors clustered by session. `p` tests that every lag of the other market is zero;
    `sum_cross` is the sum of those coefficients: the share of the other market's move it passes on over the
    lags (with its clustered standard error).
    """
    df = changes.copy()
    names = {"futures": "d_futures", "kalshi": "d_kalshi"}
    cols = []
    for l in range(1, lags + 1):
        for k, c in names.items():
            df[f"{k}_{l}"] = df.groupby("key")[c].shift(l)
            cols.append(f"{k}_{l}")
    df = df.dropna(subset=cols)
    X = np.column_stack([np.ones(len(df)), df[cols]])
    g = pd.factorize(df.session)[0]
    XtX_inv = np.linalg.inv(X.T @ X)
    rows = []
    for dep, other in [("kalshi", "futures"), ("futures", "kalshi")]:
        y = df[names[dep]].to_numpy()
        b = XtX_inv @ X.T @ y
        S = np.zeros((g.max() + 1, X.shape[1]))
        np.add.at(S, g, X * (y - X @ b)[:, None])
        V = (g.max() + 1) / g.max() * XtX_inv @ S.T @ S @ XtX_inv
        idx = [1 + cols.index(f"{other}_{l}") for l in range(1, lags + 1)]
        d, V_d = b[idx], V[np.ix_(idx, idx)]
        wald = d @ np.linalg.solve(V_d, d)
        rows.append({"question": f"Do past {other} moves predict {dep}?", "lags": lags, "changes": len(df),
                     "sessions": g.max() + 1, "wald": wald, "p": chi2_sf_even(wald, lags),
                     "sum_cross": d.sum(), "sum_cross_se": math.sqrt(V_d.sum())})
    return pd.DataFrame(rows)


@app.function(hide_code=True)
def event_response(changes, lead, follow, quantile, horizon):
    """Average path of both markets after the largest 1-minute moves of `lead`, signed so the lead's move is
    positive and scaled by its average size: 1 at minute 0 for the lead means the move itself. Paths stop at the
    end of the session, so later minutes average fewer events."""
    s = changes.reset_index(drop=True)
    size = s[lead].abs()
    events = s.index[size >= size.quantile(quantile)]
    scale = size[events].mean()
    out = {"lead": np.zeros(horizon + 1), "follow": np.zeros(horizon + 1), "n": np.zeros(horizon + 1)}
    key = s.key.to_numpy()
    for i in events:
        sign = np.sign(s.at[i, lead])
        j = np.arange(i, min(i + horizon + 1, len(s)))
        j = j[key[j] == key[i]]
        k = j - i
        out["lead"][k] += sign * s[lead].to_numpy()[j]
        out["follow"][k] += sign * s[follow].to_numpy()[j]
        out["n"][k] += 1
    n = np.where(out["n"] > 0, out["n"], np.nan)
    return pd.DataFrame({"minute": np.arange(horizon + 1),
                         "lead": np.cumsum(out["lead"] / n) / scale, "follow": np.cumsum(out["follow"] / n) / scale,
                         "events": out["n"].astype(int)}), len(events), scale


@app.cell(hide_code=True)
def intro():
    mo.md(r"""
    # 04 · Lead-lag: Kalshi vs fed funds futures, minute by minute

    Which market moves first when news about the next FOMC decision arrives? Both are put on one number, the
    **expected move at the next meeting**, in basis points, every minute:

    - **Futures.** The method of `03_fedfunds_vs_kalshi.py` (`Literature/fed_funds_implied_probabilities.pdf`),
      applied to the bid/ask mid of the contracts in force each minute instead of the daily settlement. The
      expected move is the one quantity the futures identify exactly; the bucket probabilities of notebook 03 are
      derived from it.
    - **Kalshi.** $\sum_k q_k \cdot \text{move}_k$ over the meeting's buckets (cut >25 = −50bp, …, hike >25 =
      +50bp), with $q_k$ the bid/ask mids rescaled to sum to 1.

    **Data.** 1-minute bid and ask bars for the fed funds contracts behind the next two meetings, from Bloomberg
    (`data/bloomberg/ff_intraday.xlsx`, pulled with `data/KalshiProjectFFIntraday_template.xlsx`); daily EFFR, target
    range and FOMC dates from `data/KalshiProjectFedFunds.xlsx`; Kalshi 1-minute candles from
    `data/cleaned/kalshi_candles_1m.parquet`. Bloomberg keeps about 140 business days of minute bars, so the sample
    is late March to mid-October 2026: four decided meetings and the run-up to the October one.

    **Minutes.** Bloomberg labels a bar by the minute it starts, Kalshi a candle by the minute it ends; the
    Bloomberg labels are moved one minute later so both mean "the quote at the end of this minute". Without that
    shift the futures appear to lead by exactly one minute. Only minutes when CME trades are used (Sunday 18:00 to
    Friday 17:00 ET, paused 17:00–18:00); Kalshi trades around the clock.
    """)
    return


@app.cell(hide_code=True)
def data():
    mo.stop(not INTRADAY_XLSX.exists(),
            mo.md(f"**Missing `{INTRADAY_XLSX.relative_to(ROOT)}`.** Refresh "
                  "`data/KalshiProjectFFIntraday_template.xlsx` on a Bloomberg Terminal and save it there."))
    bars = load_intraday_bars(INTRADAY_XLSX, INTRADAY_CACHE)
    fomc = load_fomc(FF_XLSX)
    meetings = fomc.meeting_date.tolist()
    daily_futures = add_starting_rate(load_futures(FF_XLSX), meetings)
    _days = pd.date_range(daily_futures.date.min(), bars.ts.max().tz_localize(None).normalize(), freq="D")
    # Weekends and the days after the workbook's last row carry the last EFFR and r0 forward.
    daily = daily_futures.set_index("date")[["effr", "r0"]].reindex(_days).ffill()
    return bars, daily, daily_futures, meetings


@app.cell(hide_code=True)
def coverage_md():
    mo.md("""
    ## 1 · The futures bars

    One row per contract: the first and last bar, the number of minutes in which the bid and the ask changed, and
    the median bid/ask spread in basis points of rate. Right: bars by hour of day (ET). The gap at 17:00 is CME's daily
    pause, which is how the time zone is identified.
    """)
    return


@app.cell(hide_code=True)
def coverage(bars):
    _cov = bars.pivot_table(index="contract_month", columns="side", values="ts", aggfunc="size")
    _cov.columns = [f"{c}_bars" for c in _cov.columns]
    _span = bars.groupby("contract_month").ts.agg(first="min", last="max")
    _q = bars.pivot_table(index=["contract_month", "ts"], columns="side", values="price", aggfunc="last")
    _q = _q.groupby(level=0).ffill().dropna()
    _spread = ((_q.ask - _q.bid) * 100).groupby(level=0).median().rename("median_spread_bp")
    _table = _span.join(_cov).join(_spread).reset_index()
    _table["first"] = _table["first"].dt.strftime("%Y-%m-%d %H:%M")
    _table["last"] = _table["last"].dt.strftime("%Y-%m-%d %H:%M")
    _hours = bars.assign(hour=(bars.ts - pd.Timedelta(minutes=1)).dt.hour).groupby("hour").size().rename("bars")
    _chart = alt.Chart(_hours.reset_index()).mark_bar(color="#2a78d6").encode(
        x=alt.X("hour:O", title="Hour of day (ET)"), y=alt.Y("bars:Q", title="Bars"),
        tooltip=["hour:O", alt.Tooltip("bars:Q", format=",")],
    ).properties(width=380, height=220, title="Bars by hour")
    mo.hstack([mo.ui.table(_table.round(2), selection=None), _chart], widths=[3, 2])
    return


@app.cell(hide_code=True)
def series_md():
    mo.md(r"""
    ## 2 · The expected move, every minute

    Futures quotes are carried forward to every CME minute; each Kalshi bucket's mid likewise (candles exist only when
    something changed), and Kalshi's expected move needs all five buckets quoted. The next meeting switches at each
    2pm decision. The starting rate $r_0$ and the EFFR already printed this month are held fixed through each CME
    session (18:00 to 17:00 ET, dated by the day it ends), so they never change between two minutes of a session.

    **Check against the daily data.** The same calculation run on the daily settlement prices of
    `KalshiProjectFedFunds.xlsx` should match the minute series at 15:00 ET, when the futures settle.
    """)
    return


@app.cell(hide_code=True)
def series(bars, daily, meetings):
    grid = cme_minutes(bars.ts.min(), bars.ts.max())
    _wide = bars.pivot_table(index="ts", columns=["contract_month", "side"], values="price", aggfunc="last")
    _wide = _wide.reindex(_wide.index.union(grid)).sort_index().ffill().reindex(grid)
    prices = pd.DataFrame({m: (_wide[(m, "bid")] + _wide[(m, "ask")]) / 2 for m in _wide.columns.levels[0]})
    # A session runs 18:00 to 17:00 ET and is dated by the day it ends.
    _session = (grid + pd.Timedelta(hours=6)).tz_localize(None).normalize()
    minute = next_meeting_move(prices, grid.tz_localize(None), _session, meetings, daily).set_index(grid)
    minute["session"] = _session

    _con = duckdb.connect()
    _con.execute((ROOT / "sql/kalshi_views.sql").read_text().replace("'data/", f"'{ROOT}/data/"))
    _k = _con.execute(f"""
        SELECT c.ticker, c.ts, (c.bid_close + c.ask_close) / 2 AS mid, m.meeting_date, m.move_bps
        FROM read_parquet('{CLEAN_CANDLES}') c JOIN markets m USING (ticker)
        WHERE m.meeting_date IN (SELECT unnest(?))
    """, [[d.date() for d in minute.meeting.unique()]]).df()
    _k["ts"] = pd.to_datetime(_k.ts).dt.tz_convert(TZ)
    minute["kalshi_bp"] = np.nan
    for _d, _g in _k.groupby("meeting_date"):
        _w = _g.pivot_table(index="ts", columns="move_bps", values="mid", aggfunc="last")
        _w = _w.reindex(_w.index.union(grid)).sort_index().ffill().reindex(grid)
        _q = _w.div(_w.sum(axis=1), axis=0)
        _e = (_q * _w.columns.to_numpy(dtype=float)).sum(axis=1, min_count=_w.shape[1])
        _rows = minute.meeting == pd.Timestamp(_d)
        minute.loc[_rows, "kalshi_bp"] = _e[_rows]
    return (minute,)


@app.cell(hide_code=True)
def validation(daily, daily_futures, meetings, minute):
    _d = daily_futures[daily_futures.date >= minute.session.min()]
    # Daily contract FFn covers the month n - 1 after the date: reshape to one column per contract month.
    _long = _d.melt(id_vars="date", value_vars=FF_COLS, var_name="n", value_name="price")
    _long["month"] = [f"{(t + pd.DateOffset(months=int(n[2:]) - 1)):%Y-%m}" for t, n in zip(_long.date, _long.n)]
    _prices = _long.pivot_table(index="date", columns="month", values="price")
    _at = _prices.index + pd.Timedelta(hours=15)
    _daily_move = next_meeting_move(_prices.reset_index(drop=True), _at, _prices.index, meetings, daily)
    _daily_move["date"] = _prices.index
    _m15 = minute[(minute.index.hour == 15) & (minute.index.minute == 0)].reset_index(drop=True)
    _m15["date"] = _m15.session
    _chk = _daily_move.merge(_m15[["date", "meeting", "futures_bp", "kalshi_bp"]], on=["date", "meeting"],
                             suffixes=("_daily", "_minute")).dropna(subset=["futures_bp_daily", "futures_bp_minute"])
    _gap = (_chk.futures_bp_minute - _chk.futures_bp_daily).abs()
    _levels = minute.dropna(subset=["futures_bp", "kalshi_bp"]).groupby("meeting").agg(
        method=("method", "first"), minutes=("futures_bp", "size"),
        futures_mean_bp=("futures_bp", "mean"), kalshi_mean_bp=("kalshi_bp", "mean"),
        futures_sd_bp=("futures_bp", "std"), kalshi_sd_bp=("kalshi_bp", "std"),
    ).reset_index()
    _levels["meeting"] = _levels.meeting.dt.strftime("%Y-%m-%d")
    mo.vstack([
        mo.md(f"**Minute series at 15:00 against daily settlements:** {len(_chk)} days, correlation "
              f"{_chk.futures_bp_minute.corr(_chk.futures_bp_daily):.4f}, median gap {_gap.median():.2f}bp, "
              f"90th percentile {_gap.quantile(0.9):.2f}bp."),
        mo.md("**Expected move per meeting** over the minutes both markets are quoted:"),
        mo.ui.table(_levels.round(2), selection=None),
    ])
    return


@app.cell(hide_code=True)
def path_controls(minute):
    _options = [d.strftime("%Y-%m-%d") for d in minute.dropna(subset=["kalshi_bp"]).meeting.unique()]
    path_meeting = mo.ui.dropdown(options=_options, value=_options[-2] if len(_options) > 1 else _options[0],
                                  label="Meeting")
    path_meeting
    return (path_meeting,)


@app.cell(hide_code=True)
def path(minute, path_meeting):
    _m = minute[minute.meeting == pd.Timestamp(path_meeting.value)][["futures_bp", "kalshi_bp"]]
    _m = _m.resample("15min").last().dropna(how="all")  # one point per 15 minutes keeps the chart light
    _long = _m.reset_index(names="time").melt(id_vars="time", var_name="source", value_name="bp")
    _long["source"] = _long.source.map({"futures_bp": "Futures", "kalshi_bp": "Kalshi"})
    _long["time"] = _long.time.dt.tz_localize(None)
    alt.Chart(_long).mark_line(strokeWidth=1.5).encode(
        x=alt.X("time:T", title=None),
        y=alt.Y("bp:Q", title="Expected move at the meeting (bp)"),
        color=alt.Color("source:N", scale=alt.Scale(domain=["Kalshi", "Futures"], range=["#2a78d6", "#eb6834"]),
                        title=None),
        tooltip=[alt.Tooltip("time:T", format="%Y-%m-%d %H:%M"), "source:N", alt.Tooltip("bp:Q", format=".2f")],
    ).properties(width="container", height=300,
                 title=f"Expected move, meeting of {path_meeting.value} (every 15 minutes, CME hours)")
    return


@app.cell(hide_code=True)
def leadlag_md():
    mo.md(r"""
    ## 3 · Who moves first

    Changes in both expected moves are taken within a CME session and a meeting (never across the 17:00 pause, a
    weekend or a decision), every 1 or 5 minutes. US hours (8:00–17:00 ET) leave out the thin overnight session.

    **Cross-correlation.** The correlation between the futures' change in one interval and Kalshi's change $k$
    intervals later. A peak at $k > 0$ means the futures move first; at $k < 0$, Kalshi does; at $k = 0$, they move
    together within the interval.

    **Granger tests.** Each market's change regressed on its own lags and the other's (10 lags at 1 minute, 6 at
    5 minutes: half an hour), standard errors clustered by session. `p` tests that the other market's lags add
    nothing; `sum_cross` is how much of the other market's move is passed on over those lags.
    """)
    return


@app.cell(hide_code=True)
def leadlag_controls():
    interval = mo.ui.dropdown(options={"1 minute": 1, "5 minutes": 5}, value="5 minutes", label="Interval")
    hours = mo.ui.dropdown(options={"All CME hours": "all", "US hours (8:00–17:00 ET)": "us"},
                           value="All CME hours", label="Hours")
    mo.hstack([interval, hours], justify="start", gap=2)
    return hours, interval


@app.cell(hide_code=True)
def changes(hours, interval, minute):
    _s = minute.dropna(subset=["futures_bp", "kalshi_bp"])
    if hours.value == "us":
        _s = _s[(_s.index.hour >= 8) & (_s.index.hour < 17)]
    _s = _s[_s.index.minute % interval.value == 0]
    _key = _s.session.dt.strftime("%Y-%m-%d") + " · " + _s.meeting.dt.strftime("%Y-%m-%d")
    changes = pd.DataFrame({
        "key": _key, "session": _s.session,
        "d_futures": _s.futures_bp.groupby(_key).diff(), "d_kalshi": _s.kalshi_bp.groupby(_key).diff(),
    }).dropna()
    # Changes need consecutive grid points: drop those that skip a gap (US-hours filter, missing quotes).
    _step = changes.index.to_series().groupby(changes.key).diff()
    changes = changes[_step.isna() | (_step == pd.Timedelta(minutes=interval.value))]
    return (changes,)


@app.cell(hide_code=True)
def leadlag(changes, interval):
    _L = LAGS[interval.value]
    _rows = []
    for _k in range(-_L, _L + 1):
        _later = changes.groupby("key").d_kalshi.shift(-_k)
        _rows.append({"lag": _k * interval.value, "corr": changes.d_futures.corr(_later),
                      "who": "Futures first" if _k > 0 else "Kalshi first" if _k < 0 else "Same interval"})
    _xc = pd.DataFrame(_rows)
    _chart = alt.Chart(_xc).mark_bar().encode(
        x=alt.X("lag:O", title="Kalshi's change, minutes after the futures' change"),
        y=alt.Y("corr:Q", title="Correlation"),
        color=alt.Color("who:N", scale=alt.Scale(domain=["Kalshi first", "Same interval", "Futures first"],
                                                 range=["#2a78d6", "#9a9994", "#eb6834"]), title=None),
        tooltip=["lag:O", alt.Tooltip("corr:Q", format=".3f"), "who:N"],
    ).properties(width="container", height=260, title="Cross-correlation of changes")
    _summary = (f"{len(changes):,} changes over {changes.session.nunique()} sessions; "
                f"{(changes.d_futures == 0).mean():.0%} of futures and {(changes.d_kalshi == 0).mean():.0%} of "
                f"Kalshi changes are zero.")
    mo.vstack([_chart, mo.md(_summary), mo.ui.table(granger(changes, _L).round(4), selection=None)])
    return


@app.cell(hide_code=True)
def event_md():
    mo.md(f"""
    ## 4 · After a large move

    The largest 1-minute moves of one market (the top {1 - EVENT_QUANTILE:.1%}, all hours), and the average path of
    both markets over the next {EVENT_HORIZON} minutes, signed so the move is upward and scaled by its average size.
    The line for the market that moved shows how much of the move sticks (below 1: part of it reverses); the other
    line shows how much of it the other market picks up, and how fast. A market that leads is one whose large moves
    the other follows.
    """)
    return


@app.cell(hide_code=True)
def events(minute):
    _s = minute.dropna(subset=["futures_bp", "kalshi_bp"])
    _key = _s.session.dt.strftime("%Y-%m-%d") + " · " + _s.meeting.dt.strftime("%Y-%m-%d")
    _c = pd.DataFrame({"key": _key, "d_futures": _s.futures_bp.groupby(_key).diff(),
                       "d_kalshi": _s.kalshi_bp.groupby(_key).diff()}).dropna()
    _charts, _notes = [], []
    for _lead, _follow, _name, _other in [("d_futures", "d_kalshi", "Futures", "Kalshi"),
                                          ("d_kalshi", "d_futures", "Kalshi", "Futures")]:
        _path, _n, _size = event_response(_c, _lead, _follow, EVENT_QUANTILE, EVENT_HORIZON)
        _long = _path.melt(id_vars=["minute", "events"], value_vars=["lead", "follow"], var_name="line",
                           value_name="share")
        _long["line"] = _long.line.map({"lead": f"{_name} (moved)", "follow": f"{_other} (follows)"})
        _charts.append(alt.Chart(_long).mark_line(strokeWidth=2).encode(
            x=alt.X("minute:Q", title="Minutes after the move"),
            y=alt.Y("share:Q", title="Cumulative change / size of the move"),
            color=alt.Color("line:N", scale=alt.Scale(
                domain=[f"{_name} (moved)", f"{_other} (follows)"],
                range=["#eb6834", "#2a78d6"] if _name == "Futures" else ["#2a78d6", "#eb6834"]), title=None),
            tooltip=["minute:Q", "line:N", alt.Tooltip("share:Q", format=".2f"), "events:Q"],
        ).properties(width=380, height=280, title=f"After a large {_name.lower()} move"))
        _at = _path.set_index("minute")
        _notes.append(f"**Large {_name.lower()} moves:** {_n} events, average size {_size:.2f}bp. After 15 minutes "
                      f"{_name.lower()} keeps {_at.lead[15]:.0%} of the move and {_other.lower()} has moved "
                      f"{_at.follow[15]:.0%} of it; after {EVENT_HORIZON} minutes {_at.lead[EVENT_HORIZON]:.0%} "
                      f"and {_at.follow[EVENT_HORIZON]:.0%}.")
    mo.vstack([mo.hstack(_charts, widths=[1, 1]), *[mo.md(n) for n in _notes]])
    return


@app.cell(hide_code=True)
def caveats():
    mo.md("""
    ## Caveats

    - **Short sample.** About 140 trading days and five meetings, the most Bloomberg keeps at one-minute resolution.
      Longer history needs a tick-data vendor (CME DataMine, Databento, LSEG Tick History).
    - **Mostly zeros.** In most minutes neither market's quote changes, so correlations are driven by the minutes in
      which something happens; 5-minute changes are the main specification for that reason.
    - **Tick sizes differ.** One futures tick moves the expected move by about 0.25–0.5bp, one Kalshi cent on a 25bp
      bucket by 0.25bp, so small moves are lumpy in both.
    - **Quotes, not trades.** Both series are bid/ask mids. A wide Kalshi book makes its mid move without anyone
      trading; the futures book is far deeper.
    - **Daily inputs.** $r_0$ and the realized EFFR come from `KalshiProjectFedFunds.xlsx` and are carried forward past
      its last row. They are constant within a session, so they do not affect the changes used in sections 3 and 4.
    """)
    return


if __name__ == "__main__":
    app.run()
