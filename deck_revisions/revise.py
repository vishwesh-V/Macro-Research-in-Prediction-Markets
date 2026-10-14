# Revise Prediction_Markets_as_Macro_Data.pptx for the instructor's "hold the audience's hand" feedback.
import sys
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches

SRC, OUT, LEADLAG_PNG = sys.argv[1:4]
E = 914400
prs = Presentation(SRC)
S = list(prs.slides)  # original order, 0-based


# ---------- helpers ----------
def shapes_named(slide, name):
    return [sh for sh in slide.shapes if sh.name == name]


def header(slide, which):
    """Kicker / title / subtitle all share the name 'Text 0'; tell them apart by height on the slide."""
    tops = {"kicker": 0.38, "title": 0.7, "subtitle": 1.34}
    for sh in shapes_named(slide, "Text 0"):
        if abs(sh.top / E - tops[which]) < 0.05:
            return sh
    raise KeyError(which)


def one(slide, name):
    found = shapes_named(slide, name)
    assert len(found) == 1, (name, len(found))
    return found[0]


def set_para(shape, text, para=0):
    """Replace a paragraph's text, keeping the first run's formatting."""
    p = shape.text_frame.paragraphs[para]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def place(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = Inches(x)
    if y is not None: shape.top = Inches(y)
    if w is not None: shape.width = Inches(w)
    if h is not None: shape.height = Inches(h)


def _next_id(slide):
    return max(int(e.get("id")) for e in slide.shapes._spTree.iter() if e.tag.endswith("}cNvPr")) + 1


def _run(text, sz, b=False, i=False, color="tx1"):
    return (f'<a:r><a:rPr lang="en-US" sz="{int(sz * 100)}"{" b=\"1\"" if b else ""}{" i=\"1\"" if i else ""} dirty="0">'
            f'<a:solidFill><a:schemeClr val="{color}"/></a:solidFill></a:rPr><a:t>{escape(text)}</a:t></a:r>')


def add_box(slide, name, x, y, w, h, paras, fill=None, line=None, geom="rect", anchor="t", inset=0.0, algn="l",
            space_after=0):
    """paras: list of paragraphs, each a list of (text, {sz, b, i, color}) runs."""
    ins = int(inset * E)
    fill_xml = f'<a:solidFill><a:schemeClr val="{fill}"/></a:solidFill>' if fill else "<a:noFill/>"
    line_xml = (f'<a:ln w="12700"><a:solidFill><a:schemeClr val="{line}"/></a:solidFill></a:ln>' if line
                else "<a:ln><a:noFill/></a:ln>")
    av = '<a:avLst><a:gd name="adj" fmla="val 50000"/></a:avLst>' if geom == "pill" else "<a:avLst/>"
    prst = "roundRect" if geom == "pill" else geom
    body = "".join(
        f'<a:p><a:pPr marL="0" indent="0" algn="{algn}"><a:spcAft><a:spcPts val="{space_after * 100}"/></a:spcAft>'
        f'<a:buNone/></a:pPr>{"".join(_run(t, **o) for t, o in runs)}</a:p>' for runs in paras)
    xml = (f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           f'<p:nvSpPr><p:cNvPr id="{_next_id(slide)}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
           f'<p:spPr><a:xfrm><a:off x="{int(x * E)}" y="{int(y * E)}"/><a:ext cx="{int(w * E)}" cy="{int(h * E)}"/>'
           f'</a:xfrm><a:prstGeom prst="{prst}">{av}</a:prstGeom>{fill_xml}{line_xml}</p:spPr>'
           f'<p:txBody><a:bodyPr wrap="square" lIns="{ins}" tIns="{ins}" rIns="{ins}" bIns="{ins}" rtlCol="0" '
           f'anchor="{anchor}"/><a:lstStyle/>{body}</p:txBody></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))


PARTS = [("1", "Calibration", 1.38), ("2", "Price discovery", 1.72), ("3", "Forecasting", 1.42)]


def tracker(slide, active, kicker):
    """Top-right tag showing which of the three parts of the talk this slide belongs to."""
    k = header(slide, "kicker")
    set_para(k, kicker)
    k.width = Inches(7.2)
    gap, right = 0.08, 12.73
    x = right - sum(w for *_, w in PARTS) - gap * (len(PARTS) - 1)
    for num, label, w in PARTS:
        on = num == active
        add_box(slide, f"part-tracker-{num}", x, 0.34, w, 0.3,
                [[(f"{num}  {label}", dict(sz=10, b=True, color="bg1" if on else "tx2"))]],
                fill="accent1" if on else "bg2", geom="pill", anchor="ctr", algn="ctr")
        x += w + gap


def footer_page(slide, text):
    for sh in shapes_named(slide, "Text 1"):
        if "/ 10" in sh.text_frame.text:
            set_para(sh, text)


L = dict(sz=13, color="tx2")  # small explanatory text, as on the existing slides
LB = dict(sz=13, b=True, color="accent1")

# ---------- 1. Motivation first, then the roadmap ----------
intro, motiv = S[1], S[2]
ids = prs.slides._sldIdLst
el = ids[2]
ids.remove(el)
ids.insert(1, el)
footer_page(motiv, "1 / 10")
footer_page(intro, "2 / 10")

# ---------- 2. Roadmap: the three parts, and which slides cover each ----------
set_para(header(intro, "kicker"), "ROADMAP")
set_para(header(intro, "subtitle"), "Three research questions, taken in order: one part of the talk each")
set_para(one(intro, "Text 4"), "Kalshi’s prices are a new, real-time read on macro expectations.")
set_para(one(intro, "Text 5"), "Are they good forecasts? We answer in three parts. "
                               "The tag in the top-right corner of every slide shows which part you are in.")
for title, pill, txt, label in [("Text 12", "Shape 13", "Text 14", "Slides 3–7 · first results"),
                                ("Text 19", "Shape 20", "Text 21", "Slide 8 · started"),
                                ("Text 26", "Shape 27", "Text 28", "Slides 9–10 · for feedback")]:
    one(intro, title).width = Inches(2.4)
    for nm in (pill, txt):
        place(one(intro, nm), x=9.83, w=2.65)
    set_para(one(intro, txt), label)

# ---------- 3. Replication: what DKW is, and why the slide is here ----------
rep = S[3]
tracker(rep, "1", "PART 1 · CALIBRATION · DATA CHECK")
set_para(header(rep, "title"), "Data check: we reproduce a Fed study, with smaller errors")
set_para(header(rep, "subtitle"), "Error of the Kalshi-implied fed funds rate vs. the rate the FOMC set; 41 meetings, 2021–2026")
add_box(rep, "dkw-definition", 0.6, 1.88, 12.13, 1.12, [
    [("What is DKW?  ", LB), ("Diercks, Katz & Wright (2026), a Federal Reserve Board working paper (FEDS 2026-010). "
                             "They turn Kalshi’s FOMC contracts into a daily probability distribution for the fed funds "
                             "rate at each meeting, and find it about as accurate as fed funds futures and surveys.",
                             dict(sz=13))],
    [("Why this slide:  ", LB), ("before testing whether Kalshi’s probabilities are right, we check that our data and "
                                 "method reproduce their result. DKW use last trades; we use bid–ask quotes.",
                                 dict(sz=13))]],
        fill="bg2", inset=0.14, space_after=3)
pic = one(rep, "fig-repl")
place(pic, x=0.6, y=3.1, h=3.75, w=7.9 * 3.75 / 4.36)
place(one(rep, "stat-0"), y=3.05)
place(one(rep, "Text 6"), y=3.73)
place(one(rep, "stat-1"), y=4.45)
place(one(rep, "Text 8"), y=5.13)
set_para(one(rep, "Text 8"), "lower average error than the DKW method, 1–160 days out (17.9 vs. 20.8bp)")
place(one(rep, "Text 9"), y=5.95)
set_para(one(rep, "Text 9"), "Why quotes win: 40–45% of last trades are over a week old.")

# ---------- 4–5. Calibration ----------
cal = S[4]
tracker(cal, "1", "PART 1 · CALIBRATION")
add_box(cal, "calibration-reading", 9.0, 5.75, 3.73, 1.0, [
    [("Calibrated = on the dashed line: a 30¢ contract pays out 30% of the time. Cheap longshots pay out even less "
      "often than priced; favorites more often.", L)]])
tracker(S[5], "1", "PART 1 · CALIBRATION")

# ---------- 6. Futures: why now, FF1–FF13, WIRP ----------
fut = S[6]
tracker(fut, "1", "PART 1 · CALIBRATION · BENCHMARK")
set_para(header(fut, "title"), "Benchmark: fed funds futures give odds for every meeting")
set_para(header(fut, "subtitle"), "RQ1 also asks if Kalshi beats existing markets, so we turn futures into the same probabilities")
place(one(fut, "Text 16"), y=3.5)
place(one(fut, "Text 17"), y=4.45)
add_box(fut, "futures-definitions", 7.75, 5.2, 4.98, 1.58, [
    [("FF1–FF13:  ", LB), ("CME 30-day fed funds futures for this month (FF1) through 12 months ahead (FF13). Each "
                           "settles at 100 minus that month’s average effective fed funds rate.", dict(sz=12))],
    [("WIRP:  ", LB), ("Bloomberg’s World Interest Rate Probability screen, the standard way traders read meeting "
                       "odds off these futures. We rebuild its method (appendix A11).", dict(sz=12))]],
        fill="bg2", inset=0.12, space_after=3)

# ---------- 7. Kalshi vs futures: define the Brier score ----------
kvf = S[7]
tracker(kvf, "1", "PART 1 · CALIBRATION · KALSHI VS. FUTURES")
br = one(kvf, "Text 8")
set_para(br, "Brier = average of (p − Y)² over hold and ±25bp contracts; p = forecast probability, Y = 1 if "
             "the outcome happened. 0 is perfect; always saying 50% scores 0.25. Fewer meetings at 45–60 days.")
place(br, h=0.8)
place(one(kvf, "brier-by-horizon"), y=3.12, h=3.5)

# ---------- 8. Lead-lag: a chart an outside viewer can read ----------
ll = S[8]
tracker(ll, "2", "PART 2 · PRICE DISCOVERY")
set_para(header(ll, "subtitle"), "Who reacts first to news? Minute-by-minute quotes, Mar–Oct 2026, 140 trading days, 5 FOMC meetings")
add_box(ll, "leadlag-heading", 0.6, 1.95, 7.9, 0.35, [[("What each bar shows", dict(sz=16, b=True))]])
add_box(ll, "leadlag-definition", 0.6, 2.3, 7.9, 0.62, [
    [("Both markets imply an expected rate change at the next meeting (bp). Each bar: correlation of the futures’ "
      "change in one 5-minute window with Kalshi’s change k minutes later.", L)]])
old = one(ll, "fig-leadlag")
_, rid = ll.part.get_or_add_image_part(LEADLAG_PNG)
old._element.blipFill.blip.rEmbed = rid
place(old, x=0.6, y=2.95, h=3.85, w=7.9 / 4.0 * 3.85)
for nm, y in [("stat-0", 2.0), ("Text 6", 2.68), ("stat-1", 3.45), ("Text 8", 4.13), ("stat-2", 5.15), ("Text 10", 5.83)]:
    place(one(ll, nm), y=y)
set_para(one(ll, "Text 6"), "same 5-minute window: most co-movement is simultaneous")
set_para(one(ll, "Text 8"), "past futures moves help predict Kalshi’s next move (Granger test); the reverse is weaker (p = 0.06)")
place(one(ll, "Text 8"), h=0.85)

# ---------- 9–10. Part 3 ----------
mean = S[9]
tracker(mean, "3", "PART 3 · NEXT STEPS")
set_para(header(mean, "subtitle"), "Our futures benchmark (slide 6) puts all probability on the two outcomes next to x")
set_para(one(mean, "Text 10"), "So Kalshi’s extra weight beyond 25bp (appendix A12) is partly mechanical.")
tracker(S[10], "3", "PART 3 · FORECASTING")

# ---------- listener-review fixes ----------
set_para(header(motiv, "subtitle"), "Contracts traded per month, all Kalshi Fed-decision markets (log scale)")
t5 = one(intro, "Text 5")
set_para(t5, "Each Kalshi contract pays $1 if its outcome (e.g. “cut 25bp”) happens, so its price in ¢ reads as a "
             "probability. Are these good forecasts? We answer in three parts; the tag at the top right of every "
             "slide shows the part.")
place(t5, y=3.3)
set_para(one(intro, "Text 29"), "Use Kalshi distributions to predict policy and macro surprises, starting with "
                                "where Kalshi and options disagree.")
set_para(header(S[5], "subtitle"), "Logistic regression of the outcome on logit(price): β = 1 means calibrated, "
                                   "β > 1 means prices are too timid")
set_para(one(S[5], "Text 8"), "joint test that α = 0 and β = 1 (cluster bootstrap)")
set_para(one(fut, "Text 12"), "strip the pre-meeting days out of that month", para=1)
set_para(one(fut, "Text 15"), "in 25bp units", para=1)
set_para(header(kvf, "title"), "Kalshi tracks futures and edges them up to 30 days out")
set_para(one(kvf, "Text 6"), "y-axis: probability of a 50bp cut; the outcome was a 50bp cut")
set_para(one(ll, "Text 10"), "of a large futures move is matched by Kalshi within 15 minutes: it follows only partly")
set_para(header(mean, "kicker"), "PART 3 · NEXT STEPS · WHY WE NEED OPTIONS")
set_para(header(mean, "subtitle"), "Part 3 needs full distributions, but our futures benchmark (slide 6) only pins down the mean x")
set_para(header(S[10], "subtitle"), "Our planned ML model: read the macro state, predict where Kalshi and options disagree")
t = one(S[10], "Text 18")
set_para(t, "Our questions for you: " + t.text_frame.text)

prs.save(OUT)
print("saved", OUT)
